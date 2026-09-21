#!/usr/bin/env python3
"""Export the built deck to a six-page PDF.

    python submission/build/export_pdf.py

Two routes, tried in order, because the right one is not always installed:

1. **LibreOffice** (`soffice --headless --convert-to pdf`). This is the correct
   export: text stays text, so the PDF is selectable, searchable and sharp at
   any zoom. Used whenever LibreOffice is present.

2. **macOS QuickLook** (`qlmanage`). The same renderer Finder uses to preview a
   `.pptx`. It renders one slide per file, so the deck is split into six
   single-slide decks, each rendered at 240 dpi and assembled into a PDF whose
   pages are exactly the slide size. The output is faithful but **rasterised**:
   the text is pixels, not glyphs.

The route actually taken is printed, and written into the PDF's Producer
metadata, so nobody has to guess afterwards which one produced the file.
Neither route alters the deck; both render what `build_deck.py` produced.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "submission" / "CryptoDrishti-SIH26164-Idea-Presentation.pptx"
PDF = DECK.with_suffix(".pdf")

DPI = 240
RID = ("{http://schemas.openxmlformats.org/officeDocument/2006/"
       "relationships}id")

SOFFICE_CANDIDATES = (
    "soffice", "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    str(Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"),
)


def find_soffice() -> str | None:
    for candidate in SOFFICE_CANDIDATES:
        found = shutil.which(candidate) or (
            candidate if Path(candidate).is_file() else None)
        if found:
            return found
    return None


# --------------------------------------------------------------------------

def via_libreoffice(soffice: str, work: Path) -> bool:
    """Convert with LibreOffice. Returns True on success."""
    print(f"  route: LibreOffice ({soffice}) -- vector text")
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf",
         "--outdir", str(work), str(DECK)],
        capture_output=True, text=True, timeout=300)
    produced = work / (DECK.stem + ".pdf")
    if result.returncode != 0 or not produced.is_file():
        print(f"  LibreOffice failed: {result.stderr.strip()[:300]}",
              file=sys.stderr)
        return False
    shutil.move(str(produced), str(PDF))
    return True


def split_slides(work: Path) -> list[Path]:
    """One single-slide copy of the deck per slide, in order.

    QuickLook renders only the first slide of a presentation, so each slide
    needs its own file. Nothing is edited: each copy is the original deck with
    the other slides removed from the slide-id list.
    """
    count = len(Presentation(str(DECK)).slides)
    parts = []
    for keep in range(count):
        prs = Presentation(str(DECK))
        lst = prs.slides._sldIdLst
        for sid in [s for i, s in enumerate(list(lst)) if i != keep]:
            prs.part.drop_rel(sid.get(RID))
            lst.remove(sid)
        part = work / f"slide-{keep + 1}.pptx"
        prs.save(str(part))
        parts.append(part)
    return parts


def via_quicklook(work: Path) -> bool:
    """Render each slide with macOS QuickLook and assemble a PDF."""
    if not shutil.which("qlmanage"):
        print("  qlmanage not available", file=sys.stderr)
        return False
    print(f"  route: macOS QuickLook -- rasterised at {DPI} dpi")

    prs = Presentation(str(DECK))
    page_w = Emu(prs.slide_width).inches
    page_h = Emu(prs.slide_height).inches
    px_w = round(page_w * DPI)
    px_h = round(page_h * DPI)

    pages: list[Image.Image] = []
    for part in split_slides(work):
        subprocess.run(["qlmanage", "-t", "-s", str(px_w), "-o", str(work),
                        str(part)], capture_output=True, timeout=120)
        rendered = work / (part.name + ".png")
        if not rendered.is_file():
            print(f"  QuickLook produced nothing for {part.name}",
                  file=sys.stderr)
            return False
        im = Image.open(rendered).convert("RGB")
        # QuickLook pads the render by a few pixels. Crop from the top-left to
        # the exact slide aspect, which drops the padding and nothing else.
        im = im.crop((0, 0, min(im.width, px_w), min(im.height, px_h)))
        if im.size != (px_w, px_h):
            im = im.resize((px_w, px_h), Image.LANCZOS)
        pages.append(im)

    if len(pages) != len(prs.slides):
        print(f"  rendered {len(pages)} pages for {len(prs.slides)} slides",
              file=sys.stderr)
        return False

    pages[0].save(PDF, "PDF", resolution=float(DPI), save_all=True,
                  append_images=pages[1:],
                  producer=f"macOS QuickLook raster export at {DPI} dpi")
    return True


def main() -> int:
    if not DECK.is_file():
        print(f"deck not found: {DECK}\nrun: python submission/build/build_deck.py",
              file=sys.stderr)
        return 1

    print(f"exporting {DECK.name}")
    with tempfile.TemporaryDirectory(prefix="cd-pdf-") as tmp:
        work = Path(tmp)
        soffice = find_soffice()
        ok = via_libreoffice(soffice, work) if soffice else False
        if not ok:
            if soffice:
                print("  falling back to QuickLook")
            else:
                print("  LibreOffice not installed; using QuickLook")
            ok = via_quicklook(work)

    if not ok:
        print("\nPDF NOT PRODUCED. Neither route is available on this machine.",
              file=sys.stderr)
        print("Install LibreOffice and re-run:\n"
              "  brew install --cask libreoffice", file=sys.stderr)
        return 1

    size = PDF.stat().st_size
    print(f"\nwritten: {PDF.relative_to(ROOT)}  ({size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

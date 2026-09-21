#!/usr/bin/env python3
"""QA the built deck: text, geometry and the fields still to be filled.

    python submission/build/check_deck.py

Three checks, because three different things can go wrong without being
visible in a thumbnail:

1. **Text** -- every string on every slide, so leftover template placeholder
   text and typos are greppable.
2. **Geometry** -- any shape that runs past the slide edge, into the template's
   footer bar, or over another shape.
3. **Completeness** -- the portal fields that are still markers rather than
   facts. A non-zero count here means the deck is not submittable yet, which
   is the honest state until the team's registration details are to hand.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "submission" / "CryptoDrishti-SIH26164-Idea-Presentation.pptx"
MARKER = "FILL FROM SIH PORTAL"

# The template's furniture. Content overlapping the footer bar or leaving the
# canvas is a defect even when a low-resolution render hides it.
FOOTER_TOP = 6.95
EDGE = 0.10

LEFTOVERS = ("lorem", "ipsum", "todo", "[insert", "maximum slides",
             "your team name", "important instruction", "xxx")

# The builder prefixes every shape it creates or repositions with `cd-`.
# Everything else is the template's own furniture: its title boxes bleed above
# the canvas by design and its decorations sit on top of one another, so those
# are not ours to police.
OURS = "cd-"


def texts(shape):
    if not shape.has_text_frame:
        return []
    return [p.text for p in shape.text_frame.paragraphs if p.text.strip()]


def box(shape):
    return (Emu(shape.left).inches, Emu(shape.top).inches,
            Emu(shape.left + shape.width).inches,
            Emu(shape.top + shape.height).inches)


def overlaps(a, b) -> bool:
    return not (a[2] <= b[0] + 0.02 or b[2] <= a[0] + 0.02
                or a[3] <= b[1] + 0.02 or b[3] <= a[1] + 0.02)


def main() -> int:
    prs = Presentation(str(DECK))
    width = Emu(prs.slide_width).inches
    height = Emu(prs.slide_height).inches
    problems: list[str] = []
    markers = 0

    print(f"{DECK.name}: {len(prs.slides)} slides, {width:.2f} x {height:.2f} in\n")
    if len(prs.slides) != 6:
        problems.append(f"slide count is {len(prs.slides)}, must be exactly 6")

    for number, slide in enumerate(prs.slides, 1):
        print(f"--- SLIDE {number} " + "-" * 56)
        content = []
        for shape in slide.shapes:
            for line in texts(shape):
                content.append(line)
                markers += line.count(MARKER)
                low = line.lower()
                for bad in LEFTOVERS:
                    if bad in low:
                        problems.append(f"slide {number}: leftover text {bad!r} in {line[:60]!r}")
        for line in content:
            print(f"    {line}")

        flagged = [s for s in slide.shapes if s.name.startswith(OURS)]
        for shape in flagged:
            left, top, right, bottom = box(shape)
            if left < -EDGE or top < -EDGE or right > width + EDGE or bottom > height + EDGE:
                problems.append(
                    f"slide {number}: {shape.name} leaves the canvas "
                    f"({left:.2f},{top:.2f})-({right:.2f},{bottom:.2f})")
            if shape in flagged and bottom > FOOTER_TOP:
                problems.append(
                    f"slide {number}: {shape.name} runs into the footer bar "
                    f"(bottom {bottom:.2f} in)")
        for i, a in enumerate(flagged):
            for b in flagged[i + 1:]:
                if overlaps(box(a), box(b)):
                    problems.append(
                        f"slide {number}: {a.name} overlaps {b.name}")
        print()

    print("=" * 70)
    if problems:
        print(f"{len(problems)} layout problem(s):")
        for line in problems:
            print(f"  ! {line}")
    else:
        print("layout: no shape leaves the canvas, hits the footer or overlaps another")

    print(f"\nportal fields still unfilled: {markers} occurrence(s) of "
          f"{MARKER!r}")
    if markers:
        print("  -> the deck is NOT ready to export as the final submission PDF.")
        print("     Fill Team ID, Team Name, Theme and the verbatim problem")
        print("     statement title in submission/build/build_deck.py, rebuild,")
        print("     and re-run this check.")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
